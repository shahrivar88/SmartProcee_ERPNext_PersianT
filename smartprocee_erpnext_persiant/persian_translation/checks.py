import hashlib
import shutil
import tempfile
from pathlib import Path

import frappe
from frappe.translate import clear_cache, get_all_translations
from frappe.utils import cint

from smartprocee_erpnext_persiant.install import after_install, restore_displaced_translations
from smartprocee_erpnext_persiant.persian_translation import operations, ownership
from smartprocee_erpnext_persiant.persian_translation.catalogue import (
	catalogue_map,
	merge_overrides,
	read_catalogue,
	write_catalogue_atomic,
)
from smartprocee_erpnext_persiant.persian_translation.excel_io import build_workbook, classify_import_rows, read_workbook
from smartprocee_erpnext_persiant.persian_translation.publish_checks import check_publish_gate, third_party_app
from smartprocee_erpnext_persiant.persian_translation.release_check import audit_catalogue
from smartprocee_erpnext_persiant.persian_translation.scan import run_scan
from smartprocee_erpnext_persiant.persian_translation.service import (
	apply_import,
	catalogue_path,
	effective_translation,
	export_excel,
	get_rows,
	preview_import,
	publish_catalogue,
	save_translations,
	translation_identity,
)
from smartprocee_erpnext_persiant.persian_translation.validate import translation_error

MARKER = "__PTM_TEST__"
OPT_IN = "persianizer_allow_destructive_checks"


def require_test_site():
	"""Refuse to run on a site that was not explicitly opted in as a test site.

	The checks overwrite and delete real Persian rows, restore snapshots and rerun
	the install hooks, so a site needs two flags in its `site_config.json`: the
	dedicated `persianizer_allow_destructive_checks` and Frappe's own
	`allow_tests`, which `bench run-tests` needs as well. `developer_mode` is no
	signal: production sites run with it too.
	"""
	missing = [key for key in (OPT_IN, "allow_tests") if not cint(frappe.conf.get(key))]
	if missing:
		frappe.throw(
			"checks.run is a destructive test harness and only runs on a test site. "
			f"This site is missing: {', '.join(missing)}. Nothing was changed.",
			frappe.PermissionError,
		)


def _hash(path: Path) -> str:
	return hashlib.sha256(path.read_bytes()).hexdigest()


def _cleanup():
	# Every key of `or_filters` is ORed, so the language belongs in `filters`;
	# otherwise `language = fa` alone matches every Persian translation.
	names = frappe.get_all(
		"Translation",
		filters={"language": "fa"},
		or_filters={
			"source_text": ["like", f"{MARKER}%"],
			"translated_text": ["like", f"{MARKER}%"],
		},
		pluck="name",
		limit_page_length=0,
	)
	for name in names:
		frappe.delete_doc("Translation", name, force=1, ignore_permissions=True)
	_forget_marker_ownership()
	backups = []
	if frappe.db.table_exists("Persian Translation Backup"):
		backups = frappe.get_all(
			"Persian Translation Backup",
			filters={"source_text": ["like", f"{MARKER}%"]},
			pluck="name",
			limit_page_length=0,
		)
	for name in backups:
		frappe.delete_doc("Persian Translation Backup", name, force=1, ignore_permissions=True)
	clear_cache()
	frappe.db.commit()


def _forget_marker_ownership():
	"""Drop the ownership claims these checks made, once their rows are gone.

	The checks also write real source strings such as "Customer Name". Ownership
	outlives a deleted row by design, so without this every run would leave a
	claim behind. A claim is removed only when its source or its last written
	value carries the marker and no Persian row exists for it any more.
	"""
	if not frappe.db.table_exists(ownership.DOCTYPE):
		return
	claims = frappe.get_all(
		ownership.DOCTYPE,
		or_filters={
			"source_text": ["like", f"{MARKER}%"],
			"last_value": ["like", f"{MARKER}%"],
		},
		fields=["name", "source_text", "context"],
		limit_page_length=0,
	)
	if not claims:
		return
	# A row saved from Desk may store its empty context as NULL.
	live = {
		(row.source_text, row.context or "")
		for row in frappe.get_all(
			"Translation",
			filters={"language": "fa", "source_text": ["in", list({claim.source_text for claim in claims})]},
			fields=["source_text", "context"],
			limit_page_length=0,
		)
	}
	for claim in claims:
		if (claim.source_text, claim.context or "") not in live:
			frappe.delete_doc(ownership.DOCTYPE, claim.name, force=1, ignore_permissions=True)
	ownership.clear_cache()


def _check(condition, message):
	if not condition:
		raise AssertionError(message)
	print("PASS", message)


def run():
	require_test_site()
	path = catalogue_path()
	before = _hash(path)
	before_count = len(read_catalogue(path))
	_check(not audit_catalogue(path).failures, f"release provenance check passes on {before_count} rows")

	_check(translation_error("Hello {0}", "سلام") is not None, "missing {0} is rejected")
	_check(translation_error("Hello {0}", "سلام {0}") is None, "preserved {0} is accepted")
	_check(translation_error("Use %(name)s", "با %(name)s") is None, "named placeholder is accepted")
	_check(translation_error("Go %s", "برو") is not None, "missing %s is rejected")
	_check(
		translation_error("See <b>Qty</b>", "مقدار") is not None,
		"removed HTML tag is rejected",
	)
	_check(translation_error("See <b>Qty</b>", "مقدار <b>Qty</b>") is None, "kept HTML tag is accepted")

	rows = read_catalogue(path)
	with tempfile.TemporaryDirectory() as tmp:
		sample = Path(tmp) / "fa.csv"
		shutil.copyfile(path, sample)
		write_catalogue_atomic(sample, read_catalogue(sample))
		_check(catalogue_map(read_catalogue(sample)) == catalogue_map(rows), "atomic rewrite keeps every catalogue value")
		key = catalogue_key_of(rows[0])
		merged, stats = merge_overrides(
			read_catalogue(sample),
			{("__PTM_TEST_NEW__", ""): "جدید", key: "تغییر آزمایشی"},
		)
		_check(stats["appended"] == 1 and stats["updated"] == 1, "merge updates one row and appends one row")
		_check(len(merged) == len(rows) + 1, "merge does not drop catalogue rows")
		shutil.copyfile(path, sample)
		original_sample = sample.read_bytes()
		try:
			write_catalogue_atomic(sample, read_catalogue(sample) + [["", "", ""]])
			failed = False
		except Exception:
			failed = True
		_check(failed and sample.read_bytes() == original_sample, "failed publish leaves the csv bytes unchanged")

	_cleanup()
	save_translations(
		[
			{"source_text": f"{MARKER} PLAIN", "context": "", "translated_text": "ساده"},
			{"source_text": f"{MARKER} CTX", "context": "Button in list view actions menu", "translated_text": "دکمه"},
			{"source_text": f"{MARKER} CTX", "context": "", "translated_text": "ساده زمینه"},
		]
	)
	frappe.db.commit()
	clear_cache()
	persian = get_all_translations("fa")
	english = get_all_translations("en")
	_check(persian.get(f"{MARKER} PLAIN") == "ساده", "fa override is visible to Frappe")
	_check(english.get(f"{MARKER} PLAIN") in (None, ""), "en dictionary does not receive the fa override")
	_check(persian.get(f"{MARKER} CTX") == "ساده زمینه", "empty context stays independent")
	_check(
		persian.get(f"{MARKER} CTX:Button in list view actions menu") == "دکمه",
		"context-specific translation stays independent",
	)
	_check(persian.get("Customer") == "مشتری", "existing Customer translation is unchanged")
	try:
		save_translations([{"source_text": "Hello {0}", "context": "", "translated_text": "بدون"}])
		rejected = False
	except Exception:
		rejected = True
	_check(rejected, "save rejects a broken placeholder")

	after_install()
	_check(
		frappe.db.exists("Translation", {"language": "fa", "source_text": f"{MARKER} PLAIN", "translated_text": "ساده"}),
		"after_install keeps administrator Translation rows",
	)

	backup = frappe.get_doc(
		{
			"doctype": "Persian Translation Backup",
			"language": "fa",
			"source_text": f"{MARKER} PLAIN",
			"translated_text": "نباید جایگزین شود",
			"context": "",
		}
	).insert(ignore_permissions=True)
	restore_displaced_translations()
	current = frappe.db.get_value(
		"Translation",
		{"language": "fa", "source_text": f"{MARKER} PLAIN"},
		"translated_text",
	)
	_check(current == "ساده", "uninstall restore does not overwrite an existing translation")
	frappe.delete_doc("Persian Translation Backup", backup.name, force=1, ignore_permissions=True)

	scanned = third_party_app() or APP_SAFE
	run_scan(apps=[scanned])
	_check(
		bool(frappe.db.exists("Persian Translation Source", {"app": scanned})),
		"scan discovers strings from an app outside the catalogue owners",
	)
	check_publish_gate(scanned, MARKER, _check)
	_cleanup()
	sidebar_label = _first_sidebar_label()
	_check(bool(sidebar_label), f"workspace sidebar extractor found {sidebar_label}")
	save_translations([{"source_text": sidebar_label, "context": "", "translated_text": f"{MARKER} منو"}])
	clear_cache()
	_check(effective_translation(sidebar_label, "") == f"{MARKER} منو", "menu label is translated by the standard _() dictionary")
	frappe.db.delete("Translation", {"language": "fa", "source_text": sidebar_label, "translated_text": f"{MARKER} منو"})
	clear_cache()
	_check(effective_translation(sidebar_label, "") != f"{MARKER} منو", "menu test translation was removed")

	identity = translation_identity(f"{MARKER} PLAIN", "")
	workbook_rows = [
		{
			"app": "test",
			"source_path": "",
			"source_text": f"{MARKER} PLAIN",
			"context": "",
			"current_translation": "ساده",
			"identity": identity,
		},
		{
			"app": "test",
			"source_path": "",
			"source_text": f"{MARKER} OTHER",
			"context": "",
			"current_translation": "",
			"identity": translation_identity(f"{MARKER} OTHER", ""),
		},
	]
	payload = build_workbook(workbook_rows)
	with tempfile.TemporaryDirectory() as tmp:
		excel = Path(tmp) / "out.xlsx"
		excel.write_bytes(payload)
		imported = read_workbook(excel)
		imported[0]["new_translation"] = "ساده عوض‌شده"
		imported[1]["new_translation"] = ""
		classified = classify_import_rows(
			imported,
			lambda source, context: "ساده" if source == f"{MARKER} PLAIN" else "",
			translation_identity,
		)
		_check(classified["changed"] == 1 and classified["unchanged"] == 1, "excel import changes only the edited row")
		_check(classified["errors"] == [], "excel import has no errors")

	_check_excel_grouping()

	broken = classify_import_rows(
		[
			{
				"source_text": "Hello {0}",
				"context": "",
				"new_translation": "خراب",
				"identity": translation_identity("Hello {0}", ""),
			}
		],
		lambda source, context: "",
		translation_identity,
	)
	_check(len(broken["errors"]) == 1 and not broken["actions"], "excel import rejects placeholder corruption")

	with tempfile.TemporaryDirectory() as tmp:
		target = Path(tmp) / "fa.csv"
		shutil.copyfile(path, target)
		one = read_catalogue(target)
		key = (one[3][0].replace("\\n", "\n"), (one[3][2] if len(one[3]) > 2 else "").replace("\\n", "\n"))
		merged, stats = merge_overrides(one, {key: "ترجمه آزمایشی انتشار", (f"{MARKER} PUBLISH", ""): "منتشرشده"})
		backup = write_catalogue_atomic(target, merged)
		reread = catalogue_map(read_catalogue(target))
		_check(reread[key] == "ترجمه آزمایشی انتشار", "publish updates one existing catalogue row")
		_check(reread[(f"{MARKER} PUBLISH", "")] == "منتشرشده", "publish appends one new catalogue row")
		_check(backup.exists() and len(read_catalogue(target)) == before_count + 1, "publish writes a backup and keeps prior rows")
		_check(stats["updated"] == 1 and stats["appended"] == 1, "publish stats count only intended rows")

	_cleanup()
	existing = frappe.db.count("Translation", {"language": "fa"})
	if existing:
		print("SKIP no-op publish: this site has its own fa Translation rows")
	else:
		no_change = publish_catalogue()
		_check(no_change.get("updated", 0) == 0 and no_change.get("appended", 0) == 0, "publish with no overrides does not rewrite fa.csv")
	_exercise_manager_apis()
	_cleanup()
	_check(frappe.db.count("Translation", {"language": "fa"}) == existing, "api checks left no test Translation rows")
	print("START full scan")
	full = run_scan()
	_check(full["strings"] > 1000, f"full scan indexed {full['strings']} strings")
	_check(_hash(path) == before, "real fa.csv hash is unchanged after the checks")
	_check(len(read_catalogue(path)) == before_count, "real fa.csv row count is unchanged")
	print("ALL CHECKS PASSED")


def _check_excel_grouping():
	from io import BytesIO

	from openpyxl import load_workbook

	long_app = "a_very_long_custom_application_name_one"
	twin_app = "a_very_long_custom_application_name_two"
	odd_app = "odd:app/name?"
	spec = [
		("erpnext", "فیلد", "Sales Invoice.customer"),
		("frappe", "فیلد", "Customer.name"),
		("erpnext", "گزارش", "Report: Stock Balance"),
		("frappe", "منو", "Workspace: Selling"),
		(long_app, "صفحه", "Page: x"),
		(twin_app, "دکمه", "Page: y"),
		(odd_app, "سایر", ""),
	]
	rows = [
		{
			"app": app,
			"kind": kind,
			"location": place,
			"source_text": f"{MARKER} {app} {kind}",
			"context": "",
			"current_translation": "",
			"identity": translation_identity(f"{MARKER} {app} {kind}", ""),
		}
		for app, kind, place in spec
	]
	book = load_workbook(BytesIO(build_workbook(rows)))
	titles = book.sheetnames
	_check(titles[:2] == ["frappe", "erpnext"] and len(titles) == 5, f"excel has one sheet per app: {titles}")
	_check(all(len(title) <= 31 and not set(title) & set(":\\/?*[]") for title in titles), "excel sheet titles are valid")
	for title in titles:
		values = list(book[title].iter_rows(min_row=2, values_only=True))
		apps = {row[0] for row in values}
		_check(len(apps) == 1, f"sheet {title} holds rows of one app")
	header = [cell.value for cell in book["erpnext"][1]]
	_check(
		header
		== [
			"برنامه",
			"نوع",
			"محل استفاده",
			"متن اصلی",
			"ترجمه فعلی",
			"ترجمه جدید",
			"وضعیت",
			"یادداشت",
			"زمینه",
			"شناسه ردیف",
		],
		f"excel header matches the manager table: {header}",
	)
	_check("ماژول" not in header, "the module column is not written to the workbook")
	locations = {row[2] for row in book["erpnext"].iter_rows(min_row=2, values_only=True)}
	_check(
		locations == {"Sales Invoice.customer", "Report: Stock Balance"},
		"the readable place stays in the محل استفاده column of one sheet",
	)
	_check(load_workbook(BytesIO(build_workbook(rows))).sheetnames == titles, "sheet titles are deterministic")


def _exercise_manager_apis():
	from io import BytesIO

	from openpyxl import load_workbook

	page = get_rows(start=0, page_length=50)
	_check(
		len(page["rows"]) == 50 and page["total"] > 1000,
		f"manager returns one page of 50 from {page['total']} indexed strings",
	)
	assign_rows = []
	start = 0
	while True:
		batch = get_rows(search="Assign To", start=start, page_length=100)
		assign_rows.extend(batch["rows"])
		start += 100
		if start >= batch["total"] or not batch["rows"]:
			break
	contexts = {(row["source_text"], row["context"]) for row in assign_rows}
	_check(
		("Assign To", "") in contexts and ("Assign To", "Button in list view actions menu") in contexts,
		"manager keeps both Assign To contexts as separate rows",
	)
	outside = third_party_app() or APP_SAFE
	_check(bool(get_rows(app=outside, page_length=10)["rows"]), "manager lists strings from an app outside the catalogue owners")

	save_translations(
		[{"source_text": "Customer Name", "context": "", "translated_text": f"{MARKER} نام مشتری"}]
	)
	frappe.db.commit()
	clear_cache()
	_check(
		get_all_translations("fa").get("Customer Name") == f"{MARKER} نام مشتری",
		"a field label override is visible to Frappe",
	)
	_check(
		get_all_translations("en").get("Customer Name") in (None, ""),
		"the field label override is absent from English",
	)

	export_excel(app=outside)
	exported = frappe.response.get("filecontent") or b""
	_check(exported[:2] == b"PK", "filtered excel export is an xlsx workbook")

	workbook_rows = [
		{
			"app": "test",
			"source_path": "",
			"source_text": f"{MARKER} IMPORT",
			"context": "",
			"current_translation": "",
			"identity": translation_identity(f"{MARKER} IMPORT", ""),
		},
		{
			"app": "test",
			"source_path": "",
			"source_text": "Customer Name",
			"context": "",
			"current_translation": f"{MARKER} نام مشتری",
			"identity": translation_identity("Customer Name", ""),
		},
	]
	book = load_workbook(BytesIO(build_workbook(workbook_rows)))
	sheet = book.active
	new_column = [cell.value for cell in sheet[1]].index("ترجمه جدید") + 1
	sheet.cell(2, new_column).value = f"{MARKER} واردشده"
	sheet.cell(3, new_column).value = ""
	payload = BytesIO()
	book.save(payload)
	from frappe.utils.file_manager import save_file

	uploaded = save_file("ptm-check.xlsx", payload.getvalue(), None, None, is_private=1)
	extra_operations = []
	extra_snapshots = []
	try:
		preview = preview_import(uploaded.file_url)
		extra_operations.append(preview["name"])
		_check(
			preview["status"] == "Queued" and preview["applied"] == 0,
			"excel preview returns a queued record before any row is read",
		)
		operations.run_operation(preview["name"])
		preview = operations.get_operation(preview["name"])
		report = preview["detail"]["report"]
		_check(
			report["new"] == 1 and report["unchanged"] == 1 and preview["error_count"] == 0,
			"excel preview counts one new row and one blank row",
		)
		_check(
			preview["planned"] == 0 and preview["checkpoint"] == 0,
			"a preview writes nothing and holds no checkpoint",
		)

		applied = apply_import(uploaded.file_url)
		extra_operations.append(applied["name"])
		_check(
			applied["status"] == "Queued" and applied["applied"] == 0,
			"the apply request returns a queued record before any row is written",
		)
		_check(
			not applied["detail"].get("undo_snapshot"),
			"a queued apply captures no undo point before its worker starts",
		)
		operations.run_operation(applied["name"])
		applied = operations.get_operation(applied["name"])
		undo = applied["detail"].get("undo_snapshot")
		if undo:
			extra_snapshots.append(undo)
		_check(applied["status"] == "Completed" and applied["applied"] == 1, "excel import writes only the changed row")
		_check(
			applied["checkpoint"] == applied["planned"] == 1,
			"the committed checkpoint agrees with the plan",
		)
		clear_cache()
		persian = get_all_translations("fa")
		_check(persian.get(f"{MARKER} IMPORT") == f"{MARKER} واردشده", "imported translation is visible to Frappe")
		_check(
			persian.get("Customer Name") == f"{MARKER} نام مشتری",
			"a blank excel cell does not delete an existing translation",
		)
		_check_restore_points()
	finally:
		for name in extra_operations:
			if frappe.db.exists(operations.OPERATION, name):
				frappe.delete_doc(operations.OPERATION, name, force=1, ignore_permissions=True)
		for name in extra_snapshots:
			if frappe.db.exists("Persian Snapshot", name):
				frappe.delete_doc("Persian Snapshot", name, force=1, ignore_on_trash=1, ignore_permissions=True)
		if uploaded and uploaded.name:
			frappe.delete_doc("File", uploaded.name, force=1, ignore_permissions=True)
		frappe.db.commit()

	frappe.set_user("Guest")
	try:
		denied = False
		try:
			get_rows()
		except frappe.PermissionError:
			denied = True
		_check(denied, "a user without System Manager cannot read the manager")
	finally:
		frappe.set_user("Administrator")


def _check_restore_points():
	"""Restore points: immutability, queued preview, targeted recovery and undo.

	Every record this creates is removed at the end, so the checks leave the site
	as they found it.
	"""
	from smartprocee_erpnext_persiant.persian_translation import operations, service, snapshots

	created = []
	made = []
	key = f"{MARKER} RESTORE"
	original = "پیش از نسخه"
	changed = "پس از نسخه"
	save_translations([{"source_text": key, "context": "", "translated_text": original}])
	clear_cache()
	try:
		baseline = snapshots.capture(snapshots.MANUAL, f"{MARKER} baseline", immutable=False)
		created.append(baseline["name"])
		doc = frappe.get_doc("Persian Snapshot", baseline["name"])
		doc.immutable = 1
		doc.save()
		created.append(doc.name)
		doc = frappe.get_doc("Persian Snapshot", doc.name)
		doc.note = "باید رد شود"
		rejected = False
		try:
			doc.save()
		except frappe.ValidationError:
			rejected = True
		_check(rejected, "an immutable restore point cannot be edited")

		save_translations([{"source_text": key, "context": "", "translated_text": changed}])
		clear_cache()
		_check(
			_service_effective(key) == changed,
			"the translation changed after the restore point was taken",
		)

		preview = service.preview_restore(baseline["name"])
		made.append(preview["name"])
		_check(
			preview["status"] == "Queued" and preview["applied"] == 0,
			"the restore preview is queued instead of running inside the request",
		)
		operations.run_operation(preview["name"])
		preview = operations.get_operation(preview["name"])
		_check(preview["status"] == "Completed", "the queued restore preview finishes")
		sample = {(row["source_text"], row["context"]): row["value"] for row in preview["detail"]["preview"]["sample"]}
		_check(sample.get((key, "")) == original, "the preview shows the recorded value, not the current one")
		_check(_service_effective(key) == changed, "the preview did not change a value")

		restore = operations.start_restore(baseline["name"], mode="values", confirm=1)
		made.append(restore["name"])
		_check(
			not restore["detail"].get("undo_snapshot"),
			"a queued restore captures no undo point before its worker starts",
		)
		operations.run_operation(restore["name"])
		restore = operations.get_operation(restore["name"])
		_check(restore["status"] == "Completed", "the targeted restore finishes")
		undo = restore["detail"].get("undo_snapshot")
		if undo:
			created.append(undo)
		_check(bool(undo), "the restore records its own undo point before it writes")
		clear_cache()
		_check(_service_effective(key) == original, "the targeted restore brought the recorded value back")
	finally:
		for name in made:
			if frappe.db.exists(operations.OPERATION, name):
				frappe.delete_doc(operations.OPERATION, name, force=1, ignore_permissions=True)
		for name in created:
			if frappe.db.exists("Persian Snapshot", name):
				frappe.delete_doc("Persian Snapshot", name, force=1, ignore_on_trash=1, ignore_permissions=True)
		for name in frappe.get_all(
			"Translation", filters={"source_text": key}, pluck="name", limit_page_length=0
		):
			frappe.delete_doc("Translation", name, force=1, ignore_permissions=True)
		print("OK restore points")


def _service_effective(source):
	from smartprocee_erpnext_persiant.persian_translation.service import effective_translation

	return effective_translation(source, "")


def seed_persistence_row():
	require_test_site()
	_cleanup()
	save_translations(
		[{"source_text": f"{MARKER} RESTART", "context": "", "translated_text": "پس از راه‌اندازی"}]
	)
	frappe.db.commit()
	print("SEEDED")


def verify_persistence_row():
	require_test_site()
	stored = frappe.db.get_value(
		"Translation",
		{"language": "fa", "source_text": f"{MARKER} RESTART"},
		"translated_text",
	)
	_check(stored == "پس از راه‌اندازی", "translation row survives process restart")
	clear_cache()
	_check(
		get_all_translations("fa").get(f"{MARKER} RESTART") == "پس از راه‌اندازی",
		"translation is visible through Frappe after restart",
	)
	_cleanup()
	remaining = frappe.db.count("Translation", {"language": "fa", "source_text": ["like", f"{MARKER}%"]})
	_check(remaining == 0, "restart test translation was removed")
	print("PERSISTENCE OK")


def catalogue_key_of(row):
	return row[0].replace("\\n", "\n"), (row[2] if len(row) > 2 else "").replace("\\n", "\n")


APP_SAFE = "smartprocee_erpnext_persiant"


def _first_sidebar_label():
	root = Path(frappe.get_app_path("erpnext"))
	from frappe.gettext.extractors import workspace_sidebar

	for path in root.rglob("*.json"):
		if path.parent.name != "workspace_sidebar":
			continue
		with path.open(encoding="utf-8") as handle:
			extracted = workspace_sidebar.extract(handle)
			items = list(extracted or [])
		for item in items:
			text = item[2] if len(item) >= 3 else ""
			if not isinstance(text, str):
				continue
			text = text.strip()
			if not text or any(token in text for token in ("{", "}", "%", "<", ">")):
				continue
			return text
	return ""
