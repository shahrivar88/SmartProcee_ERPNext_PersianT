"""The background engine behind Excel preview/apply and snapshot restore.

These tests drive a real operation record against the real database and the real
Frappe file store. The worker function is called in-process so every assertion is
deterministic; the queue itself (`frappe.enqueue` + the running RQ worker) is
exercised separately in `TestQueueing`.

What is proven here:

* a request returns a `Queued` record before any row is touched,
* rows are written in bounded batches and the checkpoint is committed with them,
* one failing row cannot undo its healthy neighbours,
* a stopped worker resumes at the last committed row without double-applying,
* the same workbook cannot be applied twice at the same time,
* counts are honest: whatever did not happen is not reported as applied.
"""

import time
from io import BytesIO
from unittest.mock import patch

import frappe
from frappe.tests import IntegrationTestCase
from frappe.utils import cint
from frappe.utils.file_manager import save_file
from openpyxl import load_workbook

from smartprocee_erpnext_persiant.persian_translation import operations, ownership, service, snapshots
from smartprocee_erpnext_persiant.persian_translation.excel_io import build_workbook
from smartprocee_erpnext_persiant.persian_translation.identity import translation_identity
from smartprocee_erpnext_persiant.persian_translation.tests.helpers import MARKER, real_queue

OPERATION = operations.OPERATION
SOURCE_COLUMN = "متن اصلی"
EDITED_COLUMN = "ترجمه جدید"
# `{0}` inside an f-string would be read as a replacement field, so it is built here.
PLACEHOLDER_SOURCE = MARKER + " THREE {0}"


def marker_row(text, current=""):
	return {
		"app": "erpnext",
		"kind": "فیلد",
		"location": "Sales Invoice.customer",
		"source_text": text,
		"context": "",
		"current_translation": current,
		"identity": translation_identity(text, ""),
	}


def edit_workbook(payload, edits):
	"""Set ترجمه جدید for the given source texts, exactly as a translator would."""
	book = load_workbook(BytesIO(payload))
	for sheet in book.worksheets:
		header = [cell.value for cell in sheet[1]]
		if SOURCE_COLUMN not in header or EDITED_COLUMN not in header:
			continue
		source_at = header.index(SOURCE_COLUMN)
		edit_at = header.index(EDITED_COLUMN)
		for row in sheet.iter_rows(min_row=2):
			if row[source_at].value in edits:
				cell = row[edit_at]
				value = edits[row[source_at].value]
				# A list gives each repeated row of that source its own value, in order.
				cell.value = value.pop(0) if isinstance(value, list) else value
				# A translator's text must stay text, even when it starts with "=".
				cell.data_type = "s"
	out = BytesIO()
	book.save(out)
	return out.getvalue()


class OperationCase(IntegrationTestCase):
	def setUp(self):
		super().setUp()
		self.files = []
		self.operations = []
		self.snapshots = []
		# Everything this test creates is removed in `cleanup_operations`, whether
		# the test created it through the service or through the engine. A queued
		# operation that survives a test run would block every later restore, so
		# the cleanup compares names instead of trusting each test to remember.
		self.existing_operations = set(frappe.get_all(OPERATION, pluck="name", limit_page_length=0))
		self.existing_snapshots = set(
			frappe.get_all(snapshots.SNAPSHOT, pluck="name", limit_page_length=0)
		)
		self.addCleanup(self.cleanup_operations)

	def cleanup_operations(self):
		frappe.db.rollback()
		for name in set(frappe.get_all(OPERATION, pluck="name", limit_page_length=0)) - self.existing_operations:
			frappe.delete_doc(OPERATION, name, force=1, ignore_permissions=True)
		for name in (
			set(frappe.get_all(snapshots.SNAPSHOT, pluck="name", limit_page_length=0))
			- self.existing_snapshots
		):
			frappe.delete_doc(
				snapshots.SNAPSHOT, name, force=1, ignore_on_trash=True, ignore_permissions=True
			)
		for url in self.files:
			name = frappe.db.get_value("File", {"file_url": url})
			if name:
				frappe.delete_doc("File", name, force=1, ignore_permissions=True)
		for field in ("source_text", "translated_text"):
			for name in frappe.get_all(
				"Translation",
				filters={"language": "fa", field: ["like", f"{MARKER}%"]},
				pluck="name",
				limit_page_length=0,
			):
				frappe.delete_doc("Translation", name, force=1, ignore_permissions=True)
		# The ownership registry is what a restore uses to decide which rows this
		# app wrote, so a leftover row would outlive the translation it describes.
		for name in frappe.get_all(
			ownership.DOCTYPE,
			filters={"source_text": ["like", f"{MARKER}%"]},
			pluck="name",
			limit_page_length=0,
		):
			frappe.delete_doc(ownership.DOCTYPE, name, force=1, ignore_permissions=True)
		service.clear_translation_caches()
		frappe.db.commit()

	def translate(self, source, value, context=""):
		"""Write one Persian value the way the manager does, so ownership is recorded."""
		service.save_translations([{"source_text": source, "context": context, "translated_text": value}])
		frappe.db.commit()

	def upload(self, rows, edits=None):
		payload = build_workbook(rows)
		if edits:
			payload = edit_workbook(payload, edits)
		file = save_file("ptm-operation-test.xlsx", payload, None, None, is_private=1)
		# The worker reads the file in another process, so it must be committed.
		frappe.db.commit()
		self.files.append(file.file_url)
		return file.file_url

	def make_operation(self, file_url, kind=operations.IMPORT_APPLY):
		name = operations._create(kind, file_url=file_url)
		self.operations.append(name)
		undo = operations.get_operation(name)["detail"].get("undo_snapshot")
		if undo:
			self.snapshots.append(undo)
		return name

	def run_now(self, file_url):
		name = self.make_operation(file_url)
		operations.run_operation(name)
		return operations.get_operation(name)

	def stored(self, source, context=""):
		return frappe.db.get_value(
			"Translation", {"language": "fa", "source_text": source, "context": context}, "translated_text"
		)


class TestApplyCounts(OperationCase):
	def test_apply_counts_what_really_happened(self):
		rows = [marker_row(f"{MARKER} ONE"), marker_row(f"{MARKER} TWO"), marker_row(PLACEHOLDER_SOURCE)]
		url = self.upload(
			rows,
			{
				f"{MARKER} ONE": "یک",
				f"{MARKER} TWO": "",  # blank cell means "leave this row alone"
				PLACEHOLDER_SOURCE: "سه",  # drops {0}, so Frappe would never accept it
			},
		)
		result = self.run_now(url)

		self.assertEqual(result["status"], "Completed")
		self.assertEqual(result["applied"], 1)
		self.assertEqual(result["unchanged"], 1)
		self.assertEqual(result["rejected"], 1)
		self.assertEqual(result["failed"], 0)
		self.assertNotIn("conflicts", result)
		self.assertEqual(result["planned"], 1)
		self.assertEqual(result["checkpoint"], 1)
		self.assertEqual(result["processed"], 1)
		self.assertFalse(result["errors_truncated"])
		self.assertEqual(result["detail"]["report"]["unchanged"], 1)

		self.assertEqual(self.stored(f"{MARKER} ONE"), "یک")
		self.assertFalse(frappe.db.exists("Translation", {"language": "fa", "source_text": PLACEHOLDER_SOURCE}))

	def test_two_rows_for_one_source_with_different_values_are_both_applied(self):
		source = f"{MARKER} TWICE"
		url = self.upload([marker_row(source), marker_row(source)], {source: ["اول", "دوم"]})

		name = self.make_operation(url, kind=operations.IMPORT_PREVIEW)
		operations.run_operation(name)
		preview = operations.get_operation(name)
		self.assertEqual(preview["status"], "Completed")
		self.assertEqual(preview["rejected"], 0)
		self.assertEqual(preview["errors"], [])
		self.assertEqual(preview["detail"]["report"]["new"], 2)

		result = self.run_now(url)
		self.assertEqual(result["status"], "Completed")
		self.assertEqual(result["applied"], 2)
		self.assertEqual(result["rejected"], 0)
		self.assertEqual(result["errors"], [])
		# Rows run in workbook order, so the later row is the value that stays.
		self.assertEqual(self.stored(source), "دوم")

	def test_a_rejected_row_is_reported_with_app_sheet_row_and_action(self):
		url = self.upload([marker_row(PLACEHOLDER_SOURCE)], {PLACEHOLDER_SOURCE: "سه"})
		result = self.run_now(url)

		self.assertEqual(len(result["errors"]), 1)
		error = result["errors"][0]
		self.assertEqual(error["code"], "placeholder_mismatch")
		self.assertEqual(error["sheet"], "erpnext")
		self.assertEqual(error["app"], "erpnext")
		self.assertEqual(error["source_text"], PLACEHOLDER_SOURCE)
		self.assertEqual(error["identity"], translation_identity(PLACEHOLDER_SOURCE, ""))
		self.assertGreaterEqual(cint(error["excel_row"]), 2)
		self.assertTrue(error["reason"])
		self.assertTrue(error["action"])
		self.assertIn("{0}", error["reason"] + error["action"] + error["source_text"])

	def test_error_workbook_is_downloadable_and_carries_every_field(self):
		url = self.upload([marker_row(PLACEHOLDER_SOURCE)], {PLACEHOLDER_SOURCE: "سه"})
		result = self.run_now(url)

		frappe.local.response = frappe._dict({})
		operations.export_errors(result["name"])
		content = frappe.response["filecontent"]
		self.assertEqual(content[:2], b"PK")
		self.assertEqual(frappe.response["type"], "binary")
		self.assertIn(result["name"], frappe.response["filename"])

		book = load_workbook(BytesIO(content))
		self.assertEqual(book.sheetnames, ["خطاها"])
		header = [cell.value for cell in book.active[1]]
		self.assertEqual(
			header,
			[
				"برنامه",
				"برگه",
				"ردیف اکسل",
				"متن اصلی",
				"ترجمهٔ ردشده",
				"زمینه",
				"علت",
				"اقدام پیشنهادی",
				"قبل",
				"بعد",
				"شناسه ردیف",
			],
		)
		row = list(next(book.active.iter_rows(min_row=2, values_only=True)))
		self.assertEqual(row[0], "erpnext")
		self.assertEqual(row[1], "erpnext")
		self.assertEqual(row[3], PLACEHOLDER_SOURCE)
		# The exact value the translator wrote, kept next to the reason that refused it.
		self.assertEqual(row[4], "سه")
		self.assertTrue(row[6])
		self.assertTrue(row[7])

	def test_a_value_that_looks_like_a_formula_is_never_a_formula(self):
		"""A rejected value is written as text: the error file must not execute it.

		`قبل` must be the value the translator wrote, not the value Frappe would
		have rewritten it into: the finding is a comparison between the two.
		"""
		source = f"{MARKER} FORMULA <a href='#'>x</a>"
		value = "=SUM(A1)&nbsp;<a href='#'>x</a>"
		url = self.upload([marker_row(source)], {source: value})
		result = self.run_now(url)

		error = result["errors"][0]
		self.assertEqual(error["code"], "html_sanitized")
		self.assertTrue(error["before"].startswith("="))
		# Stored verbatim, so the report agrees with the file instead of showing
		# what the sanitizer turned it into.
		self.assertEqual(error["source_text"], source)
		self.assertEqual(error["before"], value)
		self.assertNotEqual(error["after"], error["before"])

		frappe.local.response = frappe._dict({})
		operations.export_errors(result["name"])
		book = load_workbook(BytesIO(frappe.response["filecontent"]))
		# Columns are read from the file's own header, so a column reordering can
		# never make this test read the wrong cell and pass by accident.
		header = [cell.value for cell in book.active[1]]
		reason = book.active.cell(2, header.index("علت") + 1)
		before = book.active.cell(2, header.index("قبل") + 1)
		after = book.active.cell(2, header.index("بعد") + 1)
		self.assertEqual(before.value, value)
		self.assertEqual(before.data_type, "s")
		self.assertNotEqual(after.data_type, "f")
		self.assertNotEqual(reason.data_type, "f")

	def test_a_repeated_apply_of_the_same_file_changes_nothing(self):
		rows = [marker_row(f"{MARKER} ONE"), marker_row(f"{MARKER} TWO")]
		url = self.upload(rows, {f"{MARKER} ONE": "یک", f"{MARKER} TWO": "دو"})
		first = self.run_now(url)
		self.assertEqual(first["applied"], 2)

		second = self.run_now(url)
		self.assertEqual(second["status"], "Completed")
		self.assertEqual(second["applied"], 0)
		self.assertEqual(second["unchanged"], 2)
		self.assertEqual(second["planned"], 0)

	def test_re_running_a_committed_row_never_creates_a_second_row(self):
		url = self.upload([marker_row(f"{MARKER} ONCE")], {f"{MARKER} ONCE": "یک‌بار"})
		result = self.run_now(url)
		self.assertEqual(result["applied"], 1)

		# A replayed batch (checkpoint pushed back) must update, not duplicate.
		frappe.db.set_value(
			OPERATION, result["name"], {"status": "Queued", "checkpoint": 0}, update_modified=False
		)
		frappe.db.commit()
		operations.run_operation(result["name"])
		self.assertEqual(
			frappe.db.count("Translation", {"language": "fa", "source_text": f"{MARKER} ONCE"}), 1
		)
		self.assertEqual(self.stored(f"{MARKER} ONCE"), "یک‌بار")


class TestBatching(OperationCase):
	def test_rows_are_applied_in_bounded_committed_batches(self):
		total = 120
		rows = [marker_row(f"{MARKER} BULK {index:03d}") for index in range(total)]
		edits = {f"{MARKER} BULK {index:03d}": f"مقدار {index}" for index in range(total)}
		url = self.upload(rows, edits)

		batches = []
		real_commit = operations._commit_batch
		offset = {"at": 0}

		def spy(doc, size, counts, failures, checkpoint):
			batches.append((offset["at"], size, cint(doc.checkpoint)))
			offset["at"] += size
			return real_commit(doc, size, counts, failures, checkpoint)

		with patch.object(operations, "BATCH_SIZE", 25), patch.object(operations, "_commit_batch", spy):
			result = self.run_now(url)

		self.assertEqual(batches, [(0, 25, 0), (25, 25, 25), (50, 25, 50), (75, 25, 75), (100, 20, 100)])
		self.assertEqual(result["applied"], total)
		self.assertEqual(result["checkpoint"], total)
		self.assertEqual(result["processed"], total)
		self.assertEqual(result["failed"], 0)
		self.assertEqual(result["status"], "Completed")
		self.assertEqual(
			frappe.db.count(
				"Translation", {"language": "fa", "source_text": ["like", f"{MARKER} BULK%"]}
			),
			total,
		)

	def test_one_failing_row_does_not_undo_its_healthy_neighbours(self):
		boom = f"{MARKER} BOOM"
		rows = [marker_row(f"{MARKER} GOOD A"), marker_row(boom), marker_row(f"{MARKER} GOOD B")]
		url = self.upload(
			rows, {f"{MARKER} GOOD A": "سالم الف", boom: "خراب", f"{MARKER} GOOD B": "سالم ب"}
		)
		name = self.make_operation(url)

		real_apply = operations._apply_action

		def flaky(doc, action):
			if action.get("source_text") == boom:
				raise RuntimeError("boom")
			return real_apply(doc, action)

		with patch.object(operations, "_apply_action", flaky):
			operations.run_operation(name)
		result = operations.get_operation(name)

		self.assertEqual(result["applied"], 2)
		self.assertEqual(result["failed"], 1)
		self.assertEqual(result["status"], "Partial")
		self.assertEqual(result["checkpoint"], result["planned"])
		self.assertEqual(self.stored(f"{MARKER} GOOD A"), "سالم الف")
		self.assertEqual(self.stored(f"{MARKER} GOOD B"), "سالم ب")
		self.assertIsNone(self.stored(boom))

		failure = next(error for error in result["errors"] if error["code"] == "row_failed")
		self.assertEqual(failure["source_text"], boom)
		self.assertIn("boom", failure["reason"])
		self.assertTrue(failure["action"])


class TestResume(OperationCase):
	def test_a_stopped_worker_resumes_at_the_last_committed_row(self):
		rows = [
			marker_row(f"{MARKER} STEP {index}") for index in range(3)
		]
		url = self.upload(rows, {f"{MARKER} STEP {index}": f"گام {index}" for index in range(3)})
		name = self.make_operation(url)

		calls = {"count": 0}
		real_commit = operations._commit_batch

		def dying_commit(doc, size, counts, failures, checkpoint):
			calls["count"] += 1
			if calls["count"] > 1:
				raise RuntimeError("worker stopped")
			return real_commit(doc, size, counts, failures, checkpoint)

		with patch.object(operations, "BATCH_SIZE", 1), patch.object(operations, "_commit_batch", dying_commit):
			operations.run_operation(name)

		stopped = operations.get_operation(name)
		self.assertEqual(stopped["status"], "Failed")
		self.assertEqual(stopped["checkpoint"], 1)
		self.assertEqual(stopped["applied"], 1)
		self.assertEqual(self.stored(f"{MARKER} STEP 0"), "گام 0")
		self.assertIsNone(self.stored(f"{MARKER} STEP 1"))

		with patch.object(operations, "_enqueue") as queued:
			resumed = operations.resume_operation(name)
			self.assertEqual(queued.call_count, 1)
		self.assertEqual(resumed["status"], "Queued")
		self.assertEqual(resumed["resume_count"], 1)

		operations.run_operation(name)
		resumed = operations.get_operation(name)
		self.assertEqual(resumed["status"], "Completed")
		self.assertEqual(resumed["checkpoint"], 3)
		self.assertEqual(resumed["processed"], 3)
		self.assertEqual(resumed["resume_count"], 1)
		# 1 before the stop + 2 after it: no row was applied twice.
		self.assertEqual(resumed["applied"], 3)
		for index in range(3):
			self.assertEqual(self.stored(f"{MARKER} STEP {index}"), f"گام {index}")

	def _stopped_apply(self, label, size=3):
		sources = [f"{MARKER} {label} {index}" for index in range(size)]
		url = self.upload(
			[marker_row(source) for source in sources],
			{source: f"{label} {index}" for index, source in enumerate(sources)},
		)
		name = self.make_operation(url)
		real_commit = operations._commit_batch
		calls = {"count": 0}

		def dying_commit(doc, size, counts, failures, checkpoint):
			calls["count"] += 1
			if calls["count"] > 1:
				raise RuntimeError("worker stopped")
			return real_commit(doc, size, counts, failures, checkpoint)

		with patch.object(operations, "BATCH_SIZE", 1), patch.object(operations, "_commit_batch", dying_commit):
			operations.run_operation(name)
		self.assertEqual(frappe.db.get_value(OPERATION, name, "status"), "Failed")
		self.assertIsNone(frappe.db.get_value(OPERATION, name, "lock_key"))
		return name, sources

	def test_a_resume_takes_the_writer_lease_back_and_holds_it_until_the_end(self):
		name, sources = self._stopped_apply("LEASE RESUME")

		# Another writer holds the lease: the resume is refused and changes nothing.
		other = self.make_operation(self.upload([marker_row(f"{MARKER} LEASE OTHER")], {f"{MARKER} LEASE OTHER": "دیگر"}))
		with patch.object(operations, "_enqueue") as queued, self.assertRaises(frappe.ValidationError):
			operations.resume_operation(name)
		queued.assert_not_called()
		self.assertEqual(frappe.db.get_value(OPERATION, name, "status"), "Failed")
		self.assertIsNone(frappe.db.get_value(OPERATION, name, "lock_key"))
		operations.run_operation(other)

		with patch.object(operations, "_enqueue"):
			operations.resume_operation(name)
		self.assertEqual(frappe.db.get_value(OPERATION, name, "lock_key"), operations.LEASE_KEY)

		# While the resumed writer holds the lease, neither a second resume of the
		# same plan nor a new writer is admitted.
		with patch.object(operations, "_enqueue") as queued, self.assertRaises(frappe.ValidationError):
			operations.resume_operation(name)
		queued.assert_not_called()
		with self.assertRaises(frappe.ValidationError):
			operations._create(operations.IMPORT_APPLY, file_url=self.upload([marker_row(f"{MARKER} LEASE NEW")]))
		self.assertEqual(frappe.db.count(OPERATION, {"lock_key": operations.LEASE_KEY}), 1)

		operations.run_operation(name)
		result = operations.get_operation(name)
		self.assertEqual(result["status"], "Completed")
		self.assertEqual(result["applied"], len(sources))
		self.assertIsNone(frappe.db.get_value(OPERATION, name, "lock_key"))

	def test_a_resume_that_fails_again_releases_the_lease(self):
		name, _sources = self._stopped_apply("LEASE FAIL")
		with patch.object(operations, "_enqueue"):
			operations.resume_operation(name)

		def broken(doc, size, counts, failures, checkpoint):
			raise RuntimeError("worker stopped again")

		with patch.object(operations, "_commit_batch", broken):
			operations.run_operation(name)

		self.assertEqual(frappe.db.get_value(OPERATION, name, "status"), "Failed")
		self.assertIsNone(frappe.db.get_value(OPERATION, name, "lock_key"))
		self.operations.append(self.make_operation(self.upload([marker_row(f"{MARKER} LEASE NEXT")])))

	def test_a_finished_operation_cannot_be_resumed_again(self):
		url = self.upload([marker_row(f"{MARKER} DONE")], {f"{MARKER} DONE": "تمام"})
		result = self.run_now(url)
		with self.assertRaises(frappe.ValidationError):
			operations.resume_operation(result["name"])

	def test_cancel_is_recorded_and_stops_the_next_batch(self):
		rows = [marker_row(f"{MARKER} CANCEL {index}") for index in range(4)]
		url = self.upload(rows, {f"{MARKER} CANCEL {index}": f"لغو {index}" for index in range(4)})
		name = self.make_operation(url)
		frappe.db.set_value(OPERATION, name, "cancel_requested", 1, update_modified=False)
		frappe.db.commit()

		with patch.object(operations, "BATCH_SIZE", 1):
			operations.run_operation(name)
		result = operations.get_operation(name)

		self.assertEqual(result["status"], "Cancelled")
		self.assertEqual(result["checkpoint"], 0)
		self.assertEqual(result["applied"], 0)
		self.assertIsNone(self.stored(f"{MARKER} CANCEL 0"))


class TestProtection(OperationCase):
	def test_a_second_apply_of_the_same_workbook_is_refused_while_one_is_active(self):
		url = self.upload([marker_row(f"{MARKER} DUP")], {f"{MARKER} DUP": "یک"})
		self.make_operation(url)
		with self.assertRaises(frappe.ValidationError):
			operations._create(operations.IMPORT_APPLY, file_url=url)

	def test_two_workers_cannot_claim_the_same_operation(self):
		url = self.upload([marker_row(f"{MARKER} CLAIM")], {f"{MARKER} CLAIM": "یکی"})
		name = self.make_operation(url)

		self.assertTrue(operations._claim(name))
		# A second worker finds the row `Running`, so it must do nothing at all.
		self.assertFalse(operations._claim(name))
		self.assertTrue(operations.run_operation(name).get("skipped"))
		self.assertIsNone(self.stored(f"{MARKER} CLAIM"))

		frappe.db.set_value(OPERATION, name, "status", "Queued", update_modified=False)
		frappe.db.commit()
		self.assertTrue(operations._claim(name))

	def test_a_finished_operation_is_not_claimed_again(self):
		url = self.upload([marker_row(f"{MARKER} FINISHED")], {f"{MARKER} FINISHED": "پایان"})
		result = self.run_now(url)
		self.assertTrue(operations.run_operation(result["name"]).get("skipped"))

	def test_apply_is_refused_without_a_workbook(self):
		with self.assertRaises(frappe.ValidationError):
			operations._create(operations.IMPORT_APPLY, file_url="")

	def test_a_user_without_system_manager_cannot_read_or_start_an_operation(self):
		url = self.upload([marker_row(f"{MARKER} PERM")], {f"{MARKER} PERM": "مجوز"})
		name = self.make_operation(url)
		frappe.set_user("Guest")
		try:
			with self.assertRaises(frappe.PermissionError):
				operations.get_operation(name)
			with self.assertRaises(frappe.PermissionError):
				operations.list_operations()
			with self.assertRaises(frappe.PermissionError):
				operations.start_import(url)
			with self.assertRaises(frappe.PermissionError):
				operations.cancel_operation(name)
		finally:
			frappe.set_user("Administrator")

	def test_the_apply_records_an_undo_point_before_it_starts(self):
		url = self.upload([marker_row(f"{MARKER} UNDO")], {f"{MARKER} UNDO": "بازگشت"})
		result = self.run_now(url)
		undo = result["detail"]["undo_snapshot"]
		self.assertTrue(frappe.db.exists(snapshots.SNAPSHOT, undo))
		self.assertEqual(frappe.db.get_value(snapshots.SNAPSHOT, undo, "kind"), snapshots.PRE_APPLY)
		self.assertEqual(frappe.db.get_value(snapshots.SNAPSHOT, undo, "immutable"), 0)


class TestQueueing(OperationCase):
	def test_the_request_returns_a_queued_record_before_any_row_is_touched(self):
		"""(1) quick enqueue response on a large input, and (2) submission protection."""
		rows = [marker_row(f"{MARKER} QUEUE {index:03d}") for index in range(500)]
		edits = {f"{MARKER} QUEUE {index:03d}": f"صف {index}" for index in range(500)}
		url = self.upload(rows, edits)

		started = time.time()
		with patch.object(operations, "_enqueue") as queued:
			result = service.apply_import(url)
		elapsed = time.time() - started

		self.operations.append(result["name"])
		undo = operations.get_operation(result["name"])["detail"].get("undo_snapshot")
		if undo:
			self.snapshots.append(undo)

		self.assertEqual(queued.call_count, 1)
		self.assertEqual(queued.call_args.args[0], result["name"])
		self.assertEqual(result["status"], "Queued")
		self.assertEqual(result["processed"], 0)
		self.assertEqual(result["checkpoint"], 0)
		self.assertEqual(result["applied"], 0)
		self.assertLess(elapsed, 10, "queueing must not wait for the rows")
		self.assertFalse(frappe.db.exists("Translation", {"language": "fa", "source_text": f"{MARKER} QUEUE 000"}))

	def test_the_queued_job_is_named_deterministically(self):
		url = self.upload([marker_row(f"{MARKER} JOB")], {f"{MARKER} JOB": "کار"})
		name = self.make_operation(url)
		with patch.object(frappe, "enqueue") as queued, real_queue():
			operations._enqueue(name)
			self.assertEqual(queued.call_args.kwargs["job_id"], operations._job_id(name, 0))
			self.assertEqual(queued.call_args.kwargs["queue"], operations.QUEUE)
			self.assertTrue(queued.call_args.kwargs["enqueue_after_commit"])
			self.assertEqual(queued.call_args.kwargs["operation"], name)

	def test_the_running_worker_executes_the_queued_operation(self):
		"""The real queue and the real worker, not a mock: the job must run."""
		source = f"{MARKER} WORKER"
		url = self.upload([marker_row(source)], {source: "کارگر"})
		name = operations._create(operations.IMPORT_PREVIEW, file_url=url)
		self.operations.append(name)

		with real_queue():
			operations._enqueue(name)
		frappe.db.commit()

		from frappe.utils.background_jobs import get_job_status, is_job_enqueued

		job_id = operations._job_id(name, 0)
		self.assertEqual(frappe.db.get_value(OPERATION, name, "job_id"), job_id)
		self.assertTrue(
			is_job_enqueued(job_id) or get_job_status(job_id) is not None,
			"the job must reach the queue",
		)

		status = None
		for _ in range(90):
			frappe.db.rollback()
			status = frappe.db.get_value(OPERATION, name, "status")
			if status not in operations.ACTIVE_STATUSES:
				break
			time.sleep(1)
		self.assertIn(status, operations.FINISHED_STATUSES, "the worker never finished the job")
		self.assertEqual(status, "Completed")
		result = operations.get_operation(name)
		self.assertEqual(result["total_rows"], 1)
		self.assertEqual(cint(result["checkpoint"]), 0, "a preview writes no rows")

	def test_a_second_writer_is_refused_while_one_operation_writes(self):
		first = self.upload([marker_row(f"{MARKER} A1")], {f"{MARKER} A1": "الف"})
		second = self.upload([marker_row(f"{MARKER} B1")], {f"{MARKER} B1": "ب"})
		self.make_operation(first)
		# Two different workbooks still write the same Translation rows, so the
		# second apply waits instead of racing the first one.
		with self.assertRaises(frappe.ValidationError):
			operations._create(operations.IMPORT_APPLY, file_url=second)
		# A preview writes nothing, so a running apply does not block it.
		name = operations._create(operations.IMPORT_PREVIEW, file_url=second)
		self.operations.append(name)
		self.assertTrue(frappe.db.exists(OPERATION, name))

	def test_the_database_refuses_a_second_lease_without_the_friendly_check(self):
		"""The unique column is the guarantee; the pre-check is only a message."""
		first = self.upload([marker_row(f"{MARKER} LEASE A")], {f"{MARKER} LEASE A": "الف"})
		second = self.upload([marker_row(f"{MARKER} LEASE B")], {f"{MARKER} LEASE B": "ب"})
		self.make_operation(first)
		with patch.object(operations, "_guard"):
			with self.assertRaises(frappe.ValidationError):
				operations._create(operations.IMPORT_APPLY, file_url=second)
		self.assertEqual(
			frappe.db.count(OPERATION, {"lock_key": operations.LEASE_KEY}),
			1,
			"exactly one operation may hold the writer lease",
		)
		self.assertFalse(frappe.db.exists(OPERATION, {"file_url": second}))

	def test_finishing_an_operation_releases_the_lease(self):
		url = self.upload([marker_row(f"{MARKER} RELEASE")], {f"{MARKER} RELEASE": "آزاد"})
		result = self.run_now(url)
		self.assertEqual(result["status"], "Completed")
		self.assertIsNone(frappe.db.get_value(OPERATION, result["name"], "lock_key"))
		self.assertFalse(frappe.db.get_value(OPERATION, result["name"], "lease_token"))
		# With the lease released, the next apply is admitted.
		next_url = self.upload([marker_row(f"{MARKER} RELEASE 2")], {f"{MARKER} RELEASE 2": "آزاد ۲"})
		self.operations.append(self.make_operation(next_url))


class TestLease(OperationCase):
	def test_a_worker_that_lost_its_lease_stops_writing(self):
		rows = [marker_row(f"{MARKER} FENCE {index}") for index in range(4)]
		url = self.upload(rows, {f"{MARKER} FENCE {index}": f"نرده {index}" for index in range(4)})
		name = self.make_operation(url)
		real_commit = operations._commit_batch

		def steal(doc, size, counts, failures, checkpoint):
			result = real_commit(doc, size, counts, failures, checkpoint)
			# A recovered worker has taken the operation over.
			frappe.db.set_value(OPERATION, name, "lease_token", "other-worker", update_modified=False)
			frappe.db.commit()
			return result

		with patch.object(operations, "BATCH_SIZE", 1), patch.object(operations, "_commit_batch", steal):
			operations.run_operation(name)

		result = operations.get_operation(name)
		self.assertEqual(result["applied"], 1)
		self.assertIsNone(
			self.stored(f"{MARKER} FENCE 1"),
			"a worker that no longer owns the lease may not write another row",
		)

	def test_a_dead_worker_is_recovered_and_a_live_one_is_never_taken_over(self):
		rows = [marker_row(f"{MARKER} STALE {index}") for index in range(3)]
		url = self.upload(rows, {f"{MARKER} STALE {index}": f"کهنه {index}" for index in range(3)})
		name = self.make_operation(url)
		self.assertTrue(operations._claim(name))

		# A fresh heartbeat belongs to a worker that is still alive.
		with self.assertRaises(frappe.ValidationError):
			operations.resume_operation(name)

		frappe.db.set_value(
			OPERATION,
			name,
			"heartbeat",
			frappe.utils.add_to_date(frappe.utils.now(), seconds=-(operations.STALE_SECONDS + 60)),
			update_modified=False,
		)
		frappe.db.commit()
		self.assertTrue(operations.get_operation(name)["stale"])

		with patch.object(operations, "_enqueue") as queued:
			resumed = operations.resume_operation(name)
			self.assertEqual(queued.call_count, 1)
		self.assertEqual(resumed["status"], "Queued")

		operations.run_operation(name)
		result = operations.get_operation(name)
		self.assertEqual(result["status"], "Completed")
		self.assertEqual(result["applied"], 3)
		for index in range(3):
			self.assertEqual(self.stored(f"{MARKER} STALE {index}"), f"کهنه {index}")

	def test_a_queued_operation_is_never_stale(self):
		url = self.upload([marker_row(f"{MARKER} QUEUED STALE")], {f"{MARKER} QUEUED STALE": "صف"})
		name = self.make_operation(url)
		self.assertEqual(frappe.db.get_value(OPERATION, name, "status"), "Queued")
		self.assertFalse(operations.get_operation(name)["stale"])

	def test_an_operation_the_queue_never_ran_is_recovered(self):
		"""A job that never reached the queue may not hold the lease forever.

		A worker container that was down when the operation was queued leaves a
		`Queued` row that nothing will ever close, and its lease refuses the next
		apply or restore with a message about an operation that never started.
		"""
		source = f"{MARKER} ABANDONED"
		url = self.upload([marker_row(source)], {source: "رهاشده"})
		name = self.make_operation(url)
		frappe.db.set_value(
			OPERATION,
			name,
			{
				"job_id": "",
				"modified": frappe.utils.add_to_date(
					frappe.utils.now(), seconds=-(operations.QUEUED_GRACE_SECONDS + 60)
				),
			},
			update_modified=False,
		)
		frappe.db.commit()

		self.assertTrue(operations.is_abandoned(frappe.get_doc(OPERATION, name)))
		self.assertIn(name, operations.recover_abandoned())

		self.assertEqual(frappe.db.get_value(OPERATION, name, "status"), "Failed")
		self.assertIsNone(frappe.db.get_value(OPERATION, name, "lock_key"))
		self.assertTrue(frappe.db.get_value(OPERATION, name, "summary"))

		# The lease is free again, so a new operation is admitted instead of refused.
		other = self.make_operation(self.upload([marker_row(f"{MARKER} AFTER ABANDON")]))
		self.assertEqual(frappe.db.get_value(OPERATION, other, "status"), "Queued")

	def test_a_queued_operation_whose_job_is_still_in_the_queue_is_kept(self):
		"""The young-row window is real: the row is committed before it is enqueued."""
		url = self.upload([marker_row(f"{MARKER} QUEUED YOUNG")], {f"{MARKER} QUEUED YOUNG": "جوان"})
		name = self.make_operation(url)
		doc = frappe.get_doc(OPERATION, name)
		self.assertFalse(operations.is_abandoned(doc))
		self.assertNotIn(name, operations.recover_abandoned())
		self.assertEqual(frappe.db.get_value(OPERATION, name, "status"), "Queued")


class TestPlan(OperationCase):
	def test_every_planned_import_action_is_a_typed_translation_upsert(self):
		"""The importer describes rows; the engine executes typed actions.

		The two shapes are built in different modules, so they are asserted against
		each other here: an action that reaches the engine without `type` or
		`value` would be rejected as an unknown action type at runtime.
		"""
		source = f"{MARKER} SHAPE"
		url = self.upload([marker_row(source)], {source: "شکل"})
		name = self.make_operation(url)
		doc = frappe.get_doc(OPERATION, name)

		actions, _summary = operations._import_plan(doc)
		self.assertEqual(len(actions), 1)
		action = actions[0]
		self.assertEqual(action["type"], operations.TRANSLATION_UPSERT)
		self.assertEqual(action["value"], "شکل")
		self.assertEqual(action["source_text"], source)
		self.assertEqual(action["row_index"], 1)

		# The same dict, executed by the real engine, writes the row.
		self.assertEqual(operations._apply_action(doc, action), operations.APPLIED)
		frappe.db.commit()
		self.assertEqual(self.stored(source), "شکل")

	def test_the_plan_is_in_workbook_order_so_a_batch_cannot_step_over_a_row(self):
		"""A row that is only edited comes before a row that is new.

		Grouping new rows before changed ones would put a low row number after a
		high one, and a batch boundary at the high number would then skip the row
		whose number is lower -- the plan has to stay monotonic in `row_index`.
		"""
		edited = f"{MARKER} ORDER EDITED"
		fresh = f"{MARKER} ORDER NEW"
		self.translate(edited, "پیشین")
		rows = [marker_row(edited), marker_row(fresh)]
		url = self.upload(rows, {edited: "ویرایش‌شده", fresh: "تازه"})

		with patch.object(operations, "BATCH_SIZE", 1):
			result = self.run_now(url)

		self.assertEqual(result["status"], "Completed")
		self.assertEqual(result["applied"], 2)
		self.assertEqual(self.stored(edited), "ویرایش‌شده")
		self.assertEqual(self.stored(fresh), "تازه")

	def test_a_resume_continues_in_the_plan_the_checkpoint_measured(self):
		"""A resume must not rebuild the plan against a table it has already changed."""
		rows = [marker_row(f"{MARKER} PLAN {index}") for index in range(4)]
		sources = [f"{MARKER} PLAN {index}" for index in range(4)]
		url = self.upload(rows, {source: f"نقشه {index}" for index, source in enumerate(sources)})
		name = self.make_operation(url)

		calls = {"count": 0}
		real_commit = operations._commit_batch

		def dying(doc, size, counts, failures, checkpoint):
			calls["count"] += 1
			if calls["count"] > 1:
				raise RuntimeError("worker stopped")
			return real_commit(doc, size, counts, failures, checkpoint)

		with patch.object(operations, "BATCH_SIZE", 1), patch.object(operations, "_commit_batch", dying):
			operations.run_operation(name)

		stopped = operations.get_operation(name)
		self.assertEqual(stopped["status"], "Failed")
		self.assertEqual(stopped["checkpoint"], 1)
		self.assertTrue(stopped["plan_stored"], "the first pass must store its plan")

		stored = operations._decode_blob(
			operations._detail(frappe.get_doc(OPERATION, name).detail)["plan_blob"]
		)["actions"]
		self.assertEqual([action["source_text"] for action in stored], sources)

		# Write one of the outstanding rows behind the engine's back. A rebuild
		# would now classify it as `unchanged` and drop it from the plan, which
		# would shift every later position under the checkpoint.
		self.translate(sources[2], "نقشه 2")

		with patch.object(operations, "_enqueue"):
			operations.resume_operation(name)
		reused, _summary = operations._build_plan(frappe.get_doc(OPERATION, name))
		self.assertEqual([action["source_text"] for action in reused], sources)

		operations.run_operation(name)
		result = operations.get_operation(name)
		self.assertEqual(result["status"], "Completed")
		self.assertEqual(result["applied"], 4)
		for index, source in enumerate(sources):
			self.assertEqual(self.stored(source), f"نقشه {index}")


class TestUndoPointTiming(OperationCase):
	def test_queueing_an_apply_captures_no_snapshot_at_all(self):
		"""Reading the whole translation state belongs in the worker, not the request."""
		url = self.upload([marker_row(f"{MARKER} LATE UNDO")], {f"{MARKER} LATE UNDO": "دیر"})
		before = frappe.db.count(snapshots.SNAPSHOT)

		with patch.object(operations, "_enqueue") as queued:
			queued_result = operations.start_import(url)
			self.assertEqual(queued.call_count, 1)
		self.operations.append(queued_result["name"])

		self.assertEqual(frappe.db.count(snapshots.SNAPSHOT), before, "queueing may not write a snapshot")
		self.assertIsNone(
			operations.get_operation(queued_result["name"])["detail"].get("undo_snapshot")
		)

		operations.run_operation(queued_result["name"])
		finished = operations.get_operation(queued_result["name"])
		undo = finished["detail"]["undo_snapshot"]
		self.snapshots.append(undo)
		self.assertTrue(frappe.db.exists(snapshots.SNAPSHOT, undo))
		self.assertEqual(self.stored(f"{MARKER} LATE UNDO"), "دیر")

	def test_a_resume_reuses_the_undo_point_of_the_first_pass(self):
		rows = [marker_row(f"{MARKER} ONE UNDO {index}") for index in range(2)]
		url = self.upload(rows, {f"{MARKER} ONE UNDO {index}": f"یک بازگشت {index}" for index in range(2)})
		name = self.make_operation(url)
		operations.run_operation(name)
		first = operations.get_operation(name)["detail"]["undo_snapshot"]
		self.snapshots.append(first)

		# Force another pass over the same operation.
		frappe.db.set_value(OPERATION, name, {"status": "Queued", "checkpoint": 0}, update_modified=False)
		frappe.db.commit()
		operations.run_operation(name)
		second = operations.get_operation(name)["detail"]["undo_snapshot"]
		self.assertEqual(second, first, "a second pass may not replace the undo point")


class TestDiagnostics(OperationCase):
	def test_the_error_workbook_keeps_every_row_the_display_limit_drops(self):
		rows = [marker_row(f"{MARKER} LIMIT {index} {{0}}") for index in range(3)]
		edits = {f"{MARKER} LIMIT {index} {{0}}": "سه" for index in range(3)}
		url = self.upload(rows, edits)

		with patch.object(operations, "MAX_STORED_ERRORS", 1):
			result = self.run_now(url)

		self.assertEqual(result["rejected"], 3)
		self.assertEqual(result["error_count"], 3)
		self.assertEqual(len(result["errors"]), 1, "the stored child rows are capped")
		self.assertTrue(result["errors_truncated"])
		self.assertEqual(result["errors_full_count"], 3)

		frappe.local.response = frappe._dict({})
		operations.export_errors(result["name"])
		book = load_workbook(BytesIO(frappe.response["filecontent"]))
		self.assertEqual(book.active.max_row, 4, "the export must carry every rejected row")

	def test_a_runtime_failure_is_exportable_with_its_row(self):
		boom = f"{MARKER} EXPORT FAIL"
		url = self.upload([marker_row(boom)], {boom: "خراب"})
		name = self.make_operation(url)
		real_apply = operations._apply_action

		def flaky(doc, action):
			if action.get("source_text") == boom:
				raise RuntimeError("boom")
			return real_apply(doc, action)

		with patch.object(operations, "_apply_action", flaky):
			operations.run_operation(name)

		exported = operations._all_errors(frappe.get_doc(OPERATION, name))
		self.assertEqual([row["code"] for row in exported], ["row_failed"])
		self.assertEqual(exported[0]["source_text"], boom)
		self.assertIn("boom", exported[0]["reason"])


class TestNaming(OperationCase):
	def test_operations_their_errors_and_snapshots_use_their_own_series(self):
		url = self.upload([marker_row(PLACEHOLDER_SOURCE)], {PLACEHOLDER_SOURCE: "سه"})
		result = self.run_now(url)
		self.assertRegex(result["name"], r"^PTOP-\d{5}$")
		errors = frappe.get_all(
			"Persian Translation Operation Error", {"parent": result["name"]}, pluck="name"
		)
		self.assertTrue(errors)
		for name in errors:
			self.assertRegex(name, r"^PTOE-\d{5}$")

		snapshot = snapshots.capture(snapshots.MANUAL, f"{MARKER} naming")["name"]
		self.assertRegex(snapshot, r"^PSNAP-\d{5}$")
		# The ownership registry is keyed by what it describes, not by a counter.
		self.assertEqual(frappe.get_meta(ownership.DOCTYPE).autoname, "field:key_hash")


class TestObsoleteConflictsColumn(IntegrationTestCase):
	"""The removed conflict check's column goes away through the patch, once and safely."""

	def _has_column(self):
		return "conflicts" in frappe.db.sql_list(f"SHOW COLUMNS FROM `tab{OPERATION}`")

	def test_the_field_is_gone_from_the_doctype(self):
		self.assertFalse(frappe.get_meta(OPERATION).has_field("conflicts"))

	def test_the_patch_drops_the_column_and_is_safe_to_repeat(self):
		from smartprocee_erpnext_persiant.patches.v1_2 import drop_operation_conflicts_column as patch

		if not self._has_column():
			frappe.db.sql_ddl(f"ALTER TABLE `tab{OPERATION}` ADD COLUMN `conflicts` int(11) NOT NULL DEFAULT 0")
		self.assertTrue(self._has_column())

		patch.execute()
		self.assertFalse(self._has_column())
		self.assertFalse(frappe.db.has_column(OPERATION, "conflicts"), "the cached column list is refreshed")

		patch.execute()
		self.assertFalse(self._has_column())
