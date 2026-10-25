"""Heavy Excel export is a background operation with an owner-only download."""

import frappe
from frappe.tests import IntegrationTestCase
from openpyxl import load_workbook
from io import BytesIO

from smartprocee_erpnext_persiant.persian_translation import operations, service
from smartprocee_erpnext_persiant.persian_translation.tests.helpers import MARKER

OPERATION = operations.OPERATION


class TestExcelExportJob(IntegrationTestCase):
	def setUp(self):
		super().setUp()
		self.operations = []
		self.addCleanup(self._cleanup)

	def _cleanup(self):
		for name in self.operations:
			if frappe.db.exists(OPERATION, name):
				operations._delete_operation_files(name)
				frappe.delete_doc(OPERATION, name, force=1, ignore_permissions=True)

	def test_start_export_returns_queued_then_completes_in_the_worker(self):
		queued = operations.start_export(search=MARKER, status="all")
		self.operations.append(queued["name"])
		self.assertEqual(queued["status"], "Queued")
		self.assertEqual(queued["kind"], operations.EXCEL_EXPORT)
		self.assertFalse(queued["file_url"])
		again = operations.start_export(search=MARKER, status="all")
		self.assertEqual(again["name"], queued["name"])
		operations.run_operation(queued["name"])
		done = operations.summarize(queued["name"])
		self.assertEqual(done["status"], "Completed")
		self.assertTrue(done["file_url"])
		self.assertTrue(done["file_url"].startswith("/private/files/"))
		service.download_export(queued["name"])
		content = frappe.response.get("filecontent") or b""
		self.assertEqual(content[:2], b"PK")
		book = load_workbook(BytesIO(content))
		self.assertTrue(book.sheetnames)

	def test_a_second_export_with_different_filters_is_refused_while_one_runs(self):
		queued = operations.start_export(search=MARKER, status="untranslated")
		self.operations.append(queued["name"])
		with self.assertRaises(frappe.ValidationError):
			operations.start_export(search=MARKER, status="all")

	def test_another_system_manager_cannot_download(self):
		queued = operations.start_export(search=MARKER, app="smartprocee_erpnext_persiant")
		self.operations.append(queued["name"])
		operations.run_operation(queued["name"])
		other = "ptm.export.other@example.com"
		if not frappe.db.exists("User", other):
			user = frappe.get_doc(
				{
					"doctype": "User",
					"email": other,
					"first_name": "Export",
					"last_name": "Other",
					"send_welcome_email": 0,
					"user_type": "System User",
				}
			).insert(ignore_permissions=True)
			user.add_roles("System Manager")
			self.addCleanup(lambda: frappe.delete_doc("User", other, force=1, ignore_permissions=True))
		hidden = operations.summarize(queued["name"])
		self.assertTrue(hidden["file_url"])
		frappe.set_user(other)
		try:
			payload = operations.summarize(queued["name"])
			self.assertFalse(payload["file_url"])
			with self.assertRaises(frappe.PermissionError):
				service.download_export(queued["name"])
		finally:
			frappe.set_user("Administrator")

	def test_another_system_manager_cannot_use_the_file_api(self):
		from frappe.core.doctype.file.file import has_permission as file_controller_permission
		from frappe.core.doctype.file.utils import find_file_by_url
		from smartprocee_erpnext_persiant.permissions import file_has_permission

		queued = operations.start_export(search=MARKER, app="smartprocee_erpnext_persiant")
		self.operations.append(queued["name"])
		operations.run_operation(queued["name"])
		done = operations.summarize(queued["name"])
		file_doc = frappe.get_doc("File", {"file_url": done["file_url"]})
		self.assertTrue(file_doc.is_private)
		self.assertFalse(file_doc.attached_to_doctype)
		self.assertTrue(file_controller_permission(file_doc, "read"))
		self.assertTrue(file_has_permission(file_doc, "read"))

		other = "ptm.export.fileapi@example.com"
		if not frappe.db.exists("User", other):
			user = frappe.get_doc(
				{
					"doctype": "User",
					"email": other,
					"first_name": "Export",
					"last_name": "FileAPI",
					"send_welcome_email": 0,
					"user_type": "System User",
				}
			).insert(ignore_permissions=True)
			user.add_roles("System Manager")
			self.addCleanup(lambda: frappe.delete_doc("User", other, force=1, ignore_permissions=True))
		plain = "ptm.export.plain@example.com"
		if not frappe.db.exists("User", plain):
			guest_user = frappe.get_doc(
				{
					"doctype": "User",
					"email": plain,
					"first_name": "Export",
					"last_name": "Plain",
					"send_welcome_email": 0,
					"user_type": "System User",
				}
			).insert(ignore_permissions=True)
			self.addCleanup(lambda: frappe.delete_doc("User", plain, force=1, ignore_permissions=True))

		frappe.set_user(other)
		try:
			self.assertFalse(file_controller_permission(file_doc, "read"))
			self.assertFalse(file_has_permission(file_doc, "read"))
			self.assertFalse(file_doc.is_downloadable())
			self.assertIsNone(find_file_by_url(done["file_url"]))
			self.assertFalse(frappe.has_permission("File", "read", file_doc))
			with self.assertRaises(frappe.PermissionError):
				service.download_export(queued["name"])
		finally:
			frappe.set_user("Administrator")

		frappe.set_user(plain)
		try:
			self.assertNotIn("System Manager", frappe.get_roles())
			self.assertFalse(file_doc.is_downloadable())
			with self.assertRaises(frappe.PermissionError):
				service.download_export(queued["name"])
		finally:
			frappe.set_user("Administrator")

		frappe.set_user("Guest")
		try:
			self.assertFalse(file_doc.is_downloadable())
			with self.assertRaises(frappe.PermissionError):
				service.download_export(queued["name"])
		finally:
			frappe.set_user("Administrator")

	def test_unrelated_private_files_keep_frappe_rules(self):
		from frappe.core.doctype.file.file import has_permission as file_controller_permission
		from smartprocee_erpnext_persiant.permissions import file_has_permission

		note = frappe.get_doc({"doctype": "ToDo", "description": f"{MARKER} export-file-control"}).insert()
		self.addCleanup(lambda: frappe.delete_doc("ToDo", note.name, force=1, ignore_permissions=True))
		other_file = frappe.get_doc(
			{
				"doctype": "File",
				"file_name": f"{MARKER}-unrelated.txt",
				"content": "not an export",
				"is_private": 1,
				"attached_to_doctype": "ToDo",
				"attached_to_name": note.name,
			}
		).insert()
		self.addCleanup(lambda: frappe.delete_doc("File", other_file.name, force=1, ignore_permissions=True))
		self.assertTrue(file_has_permission(other_file, "read"))
		self.assertTrue(file_controller_permission(other_file, "read"))
