"""Excel in and out: one sheet per app, columns a translator reads.

The workbook mirrors what the manager shows. Only ``ترجمه جدید`` is meant to be
edited; the other columns are context for the translator, and the two technical
columns at the end (``زمینه`` and ``شناسه ردیف``) keep the round trip safe: an
import refuses a row whose identity does not match its source text and context,
so a value can never land on the wrong string.

No build path and no module name reach the sheet. They are useful to this app,
not to the person translating.
"""

import hashlib
import re
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

from smartprocee_erpnext_persiant.persian_translation.technical import technical_reason
from smartprocee_erpnext_persiant.persian_translation.validate import check_translation

CORE_APPS = ("frappe", "erpnext")
OTHER_SHEET = "سایر"
SHEET_TITLE_LIMIT = 31
_INVALID_SHEET_CHARS = re.compile(r"[:\\/?*\[\]]")
HEADER_SCAN_ROWS = 5

# Internal key -> the header text it is written with, then the headers an older
# workbook may still carry.
COLUMNS = (
	("app", ("برنامه", "App")),
	("kind", ("نوع", "Kind")),
	("location", ("محل استفاده", "Location")),
	("source_text", ("متن اصلی", "Source Text")),
	("current_translation", ("ترجمه فعلی", "Current Persian Translation")),
	("new_translation", ("ترجمه جدید", "New Persian Translation")),
	("status", ("وضعیت", "Status")),
	("note", ("یادداشت", "Note")),
	("context", ("زمینه", "Context")),
	("identity", ("شناسه ردیف", "Identity")),
)
HEADERS = tuple(aliases[0] for _key, aliases in COLUMNS)
WIDTHS = (12, 14, 30, 45, 40, 45, 12, 30, 24, 16)
EDITED_COLUMN = "new_translation"

# Error workbook: one row per reason a workbook row was not applied. Every row
# carries the value that was rejected (`ترجمهٔ ردشده`), so the report identifies
# the exact cell and not only the reason. `قبل` and `بعد` add the HTML Frappe
# would have received and stored when the reason is about sanitizing.
ERROR_COLUMNS = (
	("app", "برنامه"),
	("sheet", "برگه"),
	("excel_row", "ردیف اکسل"),
	("source_text", "متن اصلی"),
	("translated_text", "ترجمهٔ ردشده"),
	("context", "زمینه"),
	("reason", "علت"),
	("action", "اقدام پیشنهادی"),
	("before", "قبل"),
	("after", "بعد"),
	("identity", "شناسه ردیف"),
)
ERROR_WIDTHS = (14, 22, 10, 45, 45, 24, 45, 45, 40, 40, 16)
ERROR_SHEET = "خطاها"

# The one word the status column shows, matching the manager.
STATUS_LABELS = {
	"changed": "تغییرکرده",
	"translated": "ترجمه‌شده",
	"untranslated": "ترجمه‌نشده",
}


def build_workbook(rows) -> bytes:
	"""One worksheet per app; the App column, not the sheet title, identifies the app."""
	groups = {}
	for row in rows:
		app = (row.get("app") or "").strip() or OTHER_SHEET
		groups.setdefault(app, []).append(row)
	order = sorted(groups, key=_sheet_order)
	book = Workbook()
	used_titles = set()
	if not order:
		_write_sheet(book.active, OTHER_SHEET, [])
	for index, app in enumerate(order):
		sheet = book.active if index == 0 else book.create_sheet()
		_write_sheet(sheet, _unique_sheet_title(app, used_titles), groups[app])
	payload = BytesIO()
	book.save(payload)
	return payload.getvalue()


def read_workbook(path) -> list[dict]:
	"""Every data row of every sheet, keyed by the internal column names.

	A sheet without the expected header is skipped: a workbook is allowed to carry
	an extra sheet of notes, and only a workbook with no readable sheet at all is
	an error.
	"""
	book = load_workbook(path, read_only=True, data_only=True)
	try:
		rows = []
		found_header = False
		for sheet in book.worksheets:
			sheet_rows, sheet_header = _read_sheet(sheet)
			rows.extend(sheet_rows)
			found_header = found_header or sheet_header
		if not found_header:
			raise ValueError("ستون «متن اصلی» در فایل اکسل پیدا نشد.")
		return rows
	finally:
		book.close()


def classify_import_rows(file_rows, effective_for, identity_for) -> dict:
	"""Classify an exported workbook. Blank and unchanged Persian cells are ignored.

	Only technical checks reject a row. Two rows that give different Persian
	values for the same source and context are both accepted; they are applied in
	workbook order, so the later row is the value that stays.

	Every candidate carries `row_index`: the 1-based position of the row in the
	workbook, across sheets. `actions` is returned in that same order -- monotonic
	in `row_index` -- because progress, the checkpoint and a resume are all
	measured in it: a plan that is not monotonic would let a batch boundary skip
	a row whose number is lower than the one already committed.
	"""
	unchanged = 0
	candidates = []
	errors = []
	for index, row in enumerate(file_rows, start=2):
		row_index = index - 1
		source = (row.get("source_text") or "").strip()
		context = row.get("context") or ""
		new_value = row.get(EDITED_COLUMN) or ""
		location = row.get("excel_row") or index
		if row.get("sheet"):
			location = f"{row.get('sheet')} / {location}"
		error = _row_error(row, source, location, index)
		if error:
			errors.append(error)
			continue
		expected = identity_for(source, context)
		given = (row.get("identity") or "").strip()
		if given and given != expected:
			errors.append(
				_rejected_row(
					row,
					location=location,
					index=index,
					row_index=row_index,
					source=source,
					context=context,
					identity=given,
					code="identity_mismatch",
					message="شناسه ردیف با متن اصلی و زمینه هم‌خوان نیست.",
					action="ردیف را از فایل تازه صادرشده بردارید؛ شناسه دستی قابل تغییر نیست.",
				)
			)
			continue
		value = new_value.strip()
		if not value:
			# A blank cell means "leave this row alone".
			unchanged += 1
			continue
		current = effective_for(source, context) or ""
		if value == current.strip():
			unchanged += 1
			continue
		finding = check_translation(source, value, context)
		if finding:
			errors.append(
				_rejected_row(
					row,
					location=location,
					index=index,
					row_index=row_index,
					source=source,
					context=context,
					identity=expected,
					code=finding.code,
					message=finding.reason,
					action=finding.action,
					before=finding.before,
					after=finding.after,
				)
			)
			continue
		candidates.append(
			(
				current,
				{
					"app": row.get("app") or "",
					"source_text": source,
					"context": context,
					"translated_text": value,
					"sheet": row.get("sheet") or "",
					"excel_row": row.get("excel_row") or index,
					"row_index": row_index,
					"identity": expected,
				},
			)
		)
	new_rows = []
	changed_rows = []
	for current, payload in sorted(candidates, key=lambda item: item[1]["row_index"]):
		if current.strip() and current.strip() != payload["source_text"].strip():
			changed_rows.append(payload)
		else:
			new_rows.append(payload)
	# One list in workbook order: `checkpoint` is a workbook row number and a resume
	# only continues *after* the last committed one, so the action list must be
	# monotonic in `row_index`. Grouping new rows before changed ones would put a
	# small row number after a large one and make a batch boundary skip it.
	actions = sorted(new_rows + changed_rows, key=lambda item: item["row_index"])
	return {
		"file_rows": len(file_rows),
		"unchanged": unchanged,
		"new": len(new_rows),
		"changed": len(changed_rows),
		"errors": errors,
		"actions": actions,
	}


def _rejected_row(
	row,
	*,
	location,
	index,
	row_index,
	source,
	context,
	identity,
	code,
	message,
	action,
	before="",
	after="",
) -> dict:
	"""One rejected row, in the exact shape the stored error and the workbook use.

	`translated_text` is always the cell the translator filled in, so every report
	-- the preview list, the child table and the error workbook -- names both the
	row and the value that was refused, never only a reason.
	"""
	return {
		"app": row.get("app") or "",
		"sheet": row.get("sheet") or "",
		"excel_row": row.get("excel_row") or index,
		"row_index": row_index,
		"row": location,
		"source_text": source,
		"translated_text": row.get(EDITED_COLUMN) or "",
		"context": context,
		"identity": identity,
		"code": code,
		"message": message,
		"action": action,
		"before": before,
		"after": after,
	}


def _row_error(row, source: str, location: str, index: int) -> dict | None:
	"""Reasons that are about the row itself, before its value is looked at."""
	if not source:
		return _rejected_row(
			row,
			location=location,
			index=index,
			row_index=index - 1,
			source="",
			context=row.get("context") or "",
			identity=(row.get("identity") or "").strip(),
			code="empty_source",
			message="متن اصلی خالی است.",
			action="ردیف را از فایل حذف کنید.",
		)
	return None


def build_error_workbook(errors) -> bytes:
	"""One sheet listing every row that was not applied, with its reason.

	`قبل` and `بعد` carry the exact HTML Frappe would receive and store. Every
	cell is written as text and locked to the string type, so a value that starts
	with `=` can never become a live Excel formula and no value is altered.
	"""
	book = Workbook()
	_write_error_sheet(book.active, list(errors))
	payload = BytesIO()
	book.save(payload)
	return payload.getvalue()


def _write_error_sheet(sheet, errors):
	sheet.title = ERROR_SHEET
	sheet.append([header for _key, header in ERROR_COLUMNS])
	for position, width in enumerate(ERROR_WIDTHS, start=1):
		sheet.column_dimensions[get_column_letter(position)].width = width
	for error in errors:
		sheet.append([_error_cell(key, error) for key, _header in ERROR_COLUMNS])
		for position in range(1, len(ERROR_COLUMNS) + 1):
			sheet.cell(sheet.max_row, position).data_type = "s"
	sheet.freeze_panes = "A2"
	return sheet


def _error_cell(key: str, error) -> str:
	value = error.get(key) or ""
	return value if isinstance(value, str) else str(value)


def _write_sheet(sheet, title: str, rows) -> None:
	sheet.title = title
	sheet.append(list(HEADERS))
	for position, width in enumerate(WIDTHS, start=1):
		sheet.column_dimensions[get_column_letter(position)].width = width
	for row in rows:
		sheet.append([_cell_value(key, row) for key, _aliases in COLUMNS])
		_force_text(sheet, sheet.max_row)
	sheet.freeze_panes = "A2"


def _force_text(sheet, row: int) -> None:
	"""Keep every written cell a string.

	A Persian value that happens to start with `=` must round-trip unchanged: as a
	formula cell it would read back as an empty `data_only` value and the row would
	be treated as "leave me alone".
	"""
	for position in range(1, len(COLUMNS) + 1):
		sheet.cell(row, position).data_type = "s"


def _cell_value(key: str, row) -> str:
	if key == "status":
		status = row.get("status") or ""
		return STATUS_LABELS.get(status, status)
	if key == "note":
		return technical_reason(row.get("source_text") or "")
	if key == "new_translation":
		# Pre-filled with the current value, so a plain re-import changes nothing
		# and the translator only edits the rows they actually want to change.
		return row.get("current_translation") or ""
	return row.get(key) or ""


def _read_sheet(sheet) -> tuple[list[dict], bool]:
	"""Read one sheet. Returns the data rows and whether the header was found."""
	rows = []
	header_map = None
	for offset, raw in enumerate(sheet.iter_rows(values_only=True)):
		values = [_text(cell) for cell in raw]
		if header_map is None:
			header_map = _header_map(values, offset)
			continue
		if not any(value.strip() for value in values):
			continue
		rows.append(
			{
				**{key: _value(header_map, key, values) for key, _aliases in COLUMNS},
				"sheet": sheet.title,
				"excel_row": offset + 1,
			}
		)
	return rows, header_map is not None


def _header_map(values, offset: int) -> dict[str, int] | None:
	"""Map the header row of a sheet, or ``None`` while the header is still ahead."""
	if offset >= HEADER_SCAN_ROWS:
		return None
	headers = [value.strip() for value in values]
	mapping = {
		key: headers.index(alias)
		for key, aliases in COLUMNS
		for alias in aliases
		if alias in headers
	}
	return mapping if "source_text" in mapping else None


def _value(mapping: dict[str, int], key: str, values: list[str]) -> str:
	position = mapping.get(key)
	if position is None or position >= len(values):
		return ""
	return values[position]


def _text(cell) -> str:
	if cell is None:
		return ""
	return cell if isinstance(cell, str) else str(cell)


def _sheet_order(app: str) -> tuple:
	return (app == OTHER_SHEET, CORE_APPS.index(app) if app in CORE_APPS else len(CORE_APPS), app.casefold())


def _unique_sheet_title(label: str, used: set[str]) -> str:
	"""Excel limits: 31 characters, no :\\/?*[] and no leading/trailing apostrophe.

	A name that must be shortened or already exists gets a suffix derived from the
	full label, so the same app always receives the same title.
	"""
	cleaned = _INVALID_SHEET_CHARS.sub(" ", label or "")
	cleaned = re.sub(r"\s+", " ", cleaned).strip().strip("'").strip() or OTHER_SHEET
	candidate = cleaned
	if len(cleaned) > SHEET_TITLE_LIMIT or candidate.casefold() in used:
		suffix = "~" + hashlib.sha1(label.encode()).hexdigest()[:6]
		candidate = cleaned[: SHEET_TITLE_LIMIT - len(suffix)] + suffix
	number = 2
	while candidate.casefold() in used:
		suffix = f" {number}"
		candidate = f"{cleaned[: SHEET_TITLE_LIMIT - len(suffix)]}{suffix}"
		number += 1
	used.add(candidate.casefold())
	return candidate
